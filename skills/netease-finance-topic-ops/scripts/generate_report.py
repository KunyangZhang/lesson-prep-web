#!/usr/bin/env python3
"""Always generate the daily topic report with a consolidated source list."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

BANNED_MARKS = ("-", "—", "–")
FINAL_STATUSES = {"submitted", "approved"}


def set_font(run, name="宋体", size=11, bold=False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.superscript = False
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def chinese_date(value: str, fallback: str) -> str:
    raw = (value or "").split("T")[0]
    match = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", raw)
    if not match:
        return fallback if not raw else raw
    year, month, day = match.groups()
    return f"{year}年{int(month)}月{int(day)}日"


def display_document_no(value: str) -> str:
    raw = value or ""
    match = re.fullmatch(r"(\d{4})-(\d+)", raw)
    if match:
        return f"{match.group(1)}年第{match.group(2)}号"
    return raw


def ensure_no_banned_marks(text: str, field: str) -> None:
    found = [mark for mark in BANNED_MARKS if mark in text]
    if found:
        raise ValueError(f"{field} contains banned punctuation: {''.join(found)}")


def fact_text(fact: object, field: str) -> tuple[str, str]:
    if isinstance(fact, dict):
        conclusion = str(fact.get("conclusion", "")).strip()
        data = str(fact.get("data", "")).strip()
    elif isinstance(fact, str) and "：" in fact:
        conclusion, data = (part.strip() for part in fact.split("：", 1))
    else:
        raise ValueError(f"{field} must be an object with conclusion and data")
    if not conclusion or not data:
        raise ValueError(f"{field} conclusion and data must both be nonempty")
    if not re.fullmatch(r"[\u3400-\u9fff]{2,10}", conclusion):
        raise ValueError(f"{field} conclusion must contain 2 to 10 Chinese characters")
    ensure_no_banned_marks(conclusion, f"{field}.conclusion")
    ensure_no_banned_marks(data, f"{field}.data")
    return conclusion, data


def source_text(source: dict) -> str:
    pub = chinese_date(source.get("published_at", ""), "日期未注明")
    accessed = chinese_date(source.get("accessed_at", ""), "抓取日期未注明")
    document_no = display_document_no(source.get("document_no", ""))
    doc_no = f"，{document_no}" if document_no else ""
    grade = str(source.get("grade", "C")).upper()
    return f"【{grade}】{source.get('name', '未知来源')}：《{source.get('title', '未命名原文')}》{doc_no}，{pub}，{source.get('url', '')}（抓取日期：{accessed}）"


def validate_topic(topic: dict) -> None:
    topic_id = str(topic.get("topic_id", "未命名选题"))
    title = str(topic.get("suggested_title", ""))
    if not title:
        raise ValueError(f"{topic_id} missing suggested_title")
    ensure_no_banned_marks(title, f"{topic_id}.suggested_title")
    sources = {source.get("source_id"): source for source in topic.get("sources", [])}
    sections = topic.get("sections", [])
    if len(sections) != 3:
        raise ValueError(f"{topic_id} must contain exactly three sections")
    for section_item in sections:
        ensure_no_banned_marks(section_item.get("title", ""), f"{topic_id}.section.title")
        for point in section_item.get("points", []):
            ensure_no_banned_marks(point.get("label", ""), f"{topic_id}.point.label")
            for fact_idx, fact in enumerate(point.get("facts", [])):
                fact_text(fact, f"{topic_id}.facts[{fact_idx}]")
            for source_id in point.get("source_ids", []):
                if source_id not in sources:
                    raise ValueError(f"unknown source_id {source_id} in {topic_id}")


def completion_failures(payload: dict, topic_count: int, skipped: list[str]) -> list[str]:
    gate = payload.get("search_gate") or {}
    failures = [str(item) for item in gate.get("failures", []) if str(item).strip()]
    if topic_count < 4:
        failures.append(f"成功选题不足4条：{topic_count}条")
    failures.extend(f"选题格式不合格：{reason}" for reason in skipped)
    return failures


def build(args: argparse.Namespace) -> None:
    payload = json.loads(Path(args.input).read_text(encoding="utf-8"))
    candidates = [topic for topic in payload.get("topics", []) if topic.get("status") in FINAL_STATUSES]
    candidates.sort(key=lambda topic: (-int(topic.get("score_total", sum((topic.get("scores") or {}).values()))), topic.get("topic_id", "")))
    topics = []
    skipped = []
    for topic in candidates:
        try:
            validate_topic(topic)
            topics.append(topic)
        except (KeyError, TypeError, ValueError) as exc:
            skipped.append(str(exc))
    failures = completion_failures(payload, len(topics), skipped)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Pt(54)
    section.left_margin = section.right_margin = Pt(68)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    date_text = chinese_date(payload.get("run", {}).get("search_date", ""), "日期未注明")
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(heading.add_run(date_text), size=12, bold=True)

    status = doc.add_paragraph()
    set_font(status.add_run("完成状态："), bold=True)
    set_font(status.add_run("达标" if not failures else "未达标"))
    summary = doc.add_paragraph()
    set_font(summary.add_run("实际结果："), bold=True)
    set_font(summary.add_run(f"形成{len(topics)}条成功选题"))
    if failures:
        reason = doc.add_paragraph()
        set_font(reason.add_run("未达标原因："), bold=True)
        set_font(reason.add_run("；".join(failures)))

    used_sources: dict[str, dict] = {}
    if not topics:
        empty = doc.add_paragraph()
        set_font(empty.add_run("本次未形成符合全部门槛的成功选题。"))

    for topic_no, topic in enumerate(topics, 1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(6)
        set_font(paragraph.add_run(f"选题{topic_no}：{topic['suggested_title']}"), size=12, bold=True)
        sources = {source.get("source_id"): source for source in topic.get("sources", [])}
        sections = topic.get("sections", [])
        for section_item in sections:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(3)
            set_font(paragraph.add_run(section_item.get("title", "")), bold=True)
            for point in section_item.get("points", []):
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Pt(12)
                paragraph.paragraph_format.space_after = Pt(2)
                set_font(paragraph.add_run(point.get("label", "")), bold=True)
                for fact_idx, fact in enumerate(point.get("facts", [])):
                    conclusion, data = fact_text(fact, f"{topic['topic_id']}.facts[{fact_idx}]")
                    paragraph = doc.add_paragraph()
                    paragraph.paragraph_format.left_indent = Pt(24)
                    paragraph.paragraph_format.space_after = Pt(1)
                    set_font(paragraph.add_run(f"{conclusion}："), bold=True)
                    set_font(paragraph.add_run(data))
                for source_id in point.get("source_ids", []):
                    source = sources.get(source_id)
                    used_sources.setdefault(str(source_id), source)

    sources_heading = doc.add_paragraph()
    sources_heading.paragraph_format.space_before = Pt(12)
    set_font(sources_heading.add_run("信息来源"), size=12, bold=True)
    if used_sources:
        for index, source in enumerate(used_sources.values(), 1):
            paragraph = doc.add_paragraph()
            set_font(paragraph.add_run(f"{index}. {source_text(source)}"), size=9)
    else:
        paragraph = doc.add_paragraph()
        set_font(paragraph.add_run("本次没有成功选题引用的信息来源。"), size=9)

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(json.dumps({"topics": len(topics), "sources": len(used_sources), "complete": not failures, "failures": failures, "output": str(output.resolve())}, ensure_ascii=False))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    build(args)


if __name__ == "__main__":
    main()
